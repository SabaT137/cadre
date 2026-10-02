"""Detect and fill placeholders in .docx contract templates without losing formatting.

Paragraphs are addressed by a single global index over the body (including table
cells, in document order) followed by headers and footers. The same file always
yields the same indices, so `scan` output can be fed straight into `fill`.
"""
from __future__ import annotations

import copy
import re
from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

# [●] / [•] / [ ● ], [insert relevant details], {{field_name}}, ______
INLINE_RE = re.compile(
    r"\[\s*[●•]\s*\]|\[\s*insert[^\]]*\]|\{\{\s*[\w.]+\s*\}\}|_{4,}",
    re.IGNORECASE,
)
# A paragraph that is just a label waiting for a value, e.g. "Name:" or "1- \tCNIC/Passport No:"
BLANK_FIELD_RE = re.compile(r"^\s*(?:\d+\s*-\s*)?\s*([A-Za-z][A-Za-z/ .]{0,40}):\s*$")


@dataclass
class Placeholder:
    id: int
    paragraph_idx: int
    kind: str  # "inline" | "blank_field"
    token: str
    section: str
    context: str
    start: int
    end: int


def _paragraphs(doc) -> list[Paragraph]:
    paras: list[Paragraph] = []
    for p_el in doc.element.body.iter(qn("w:p")):
        paras.append(Paragraph(p_el, doc._body))
    for section in doc.sections:
        for part in (section.header, section.footer):
            if part.is_linked_to_previous:
                continue
            for p_el in part._element.iter(qn("w:p")):
                paras.append(Paragraph(p_el, part))
    return paras


def _text(p: Paragraph) -> str:
    return "".join(r.text for r in p.runs)


def _is_heading(p: Paragraph, text: str) -> bool:
    t = text.strip()
    if not t or len(t) > 80 or INLINE_RE.search(t):
        return False
    if p.style is not None and p.style.name.lower().startswith(("heading", "title")):
        return True
    letters = [c for c in t if c.isalpha()]
    if letters and all(c.isupper() for c in letters) and len(letters) > 3:
        return True
    if t.lower().startswith("for and on behalf of"):
        return True
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and all(r.bold for r in runs) and len(t) < 40


def _context(text: str, start: int, end: int, window: int = 110) -> str:
    left = text[max(0, start - window):start]
    right = text[end:end + window]
    left = INLINE_RE.sub("[…]", left)
    right = INLINE_RE.sub("[…]", right)
    return f"{'…' if start > window else ''}{left}<<HERE>>{right}{'…' if end + window < len(text) else ''}".strip()


def scan_document(doc) -> list[Placeholder]:
    found: list[Placeholder] = []
    section = ""
    prev_texts: list[str] = []
    for idx, p in enumerate(_paragraphs(doc)):
        text = _text(p)
        if _is_heading(p, text):
            section = text.strip()
        matches = list(INLINE_RE.finditer(text))
        for m in matches:
            ctx = _context(text, m.start(), m.end())
            if len(text.strip()) <= len(m.group()) + 3 and prev_texts:
                # Bare "[●]" bullet: show what precedes it so the model knows what list it is in.
                ctx = f"(after: {prev_texts[-1][:90]}) {ctx}"
            found.append(Placeholder(len(found) + 1, idx, "inline", m.group(), section, ctx, m.start(), m.end()))
        if not matches and BLANK_FIELD_RE.match(text):
            before = " / ".join(t[:60] for t in prev_texts[-2:])
            ctx = f"(after: {before}) {text.strip()} <<HERE>>"
            found.append(Placeholder(len(found) + 1, idx, "blank_field", text.strip(), section, ctx, len(text), len(text)))
        if text.strip():
            prev_texts.append(text.strip())
    return found


def scan(path: str | Path) -> list[dict]:
    return [asdict(p) for p in scan_document(Document(str(path)))]


def _replace_span(p: Paragraph, start: int, end: int, value: str) -> None:
    """Replace characters [start, end) of the paragraph's run text, keeping the first run's formatting."""
    runs = [r for r in p.runs]
    if start == end:  # insertion after the label of a blank field
        last = next((r for r in reversed(runs) if r.text), None)
        if last is None:
            p.add_run(value)
        else:
            sep = "" if last.text.endswith((" ", "\t")) else " "
            last.text = last.text + sep + value
        return
    pos = 0
    first = True
    for run in runs:
        rt = run.text
        r_start, r_end = pos, pos + len(rt)
        pos = r_end
        if r_end <= start or r_start >= end:
            continue
        lo = max(start, r_start) - r_start
        hi = min(end, r_end) - r_start
        if first:
            run.text = rt[:lo] + value + rt[hi:]
            first = False
        else:
            run.text = rt[:lo] + rt[hi:]


def _set_paragraph_text(p: Paragraph, new_text: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(new_text)
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def fill(
    template_path: str | Path,
    output_path: str | Path,
    values: dict[int, str],
    clause_edits: list[dict] | None = None,
) -> dict:
    """Fill placeholders by id and apply clause edits.

    clause_edits items: {"paragraph_idx": int, "action": "replace"|"insert_after"|"delete", "text": str}
    Edits refer to paragraph indices of the *original* template.
    """
    doc = Document(str(template_path))
    placeholders = scan_document(doc)
    by_id = {ph.id: ph for ph in placeholders}
    values = {int(k): str(v) for k, v in values.items() if str(v).strip()}
    unknown = sorted(k for k in values if k not in by_id)

    paras = _paragraphs(doc)
    # Replace right-to-left within each paragraph so earlier offsets stay valid.
    for ph in sorted(placeholders, key=lambda x: (x.paragraph_idx, x.start), reverse=True):
        if ph.id not in values:
            continue
        p = paras[ph.paragraph_idx]
        lines = [re.sub(r"^\s*[-•*]\s*", "", ln).strip() for ln in values[ph.id].splitlines()]
        lines = [ln for ln in lines if ln]
        if ph.kind == "inline" and _text(p).strip() == ph.token and len(lines) > 1:
            # A bullet that received several lines becomes several bullets with the same formatting.
            anchor = p._p
            for extra in lines[1:]:
                new_el = copy.deepcopy(p._p)
                anchor.addnext(new_el)
                _set_paragraph_text(Paragraph(new_el, p._parent), extra)
                anchor = new_el
            _set_paragraph_text(p, lines[0])
        else:
            _replace_span(p, ph.start, ph.end, " ".join(lines) if ph.kind == "inline" else values[ph.id].strip())

    # Unused list items that are nothing but a placeholder (e.g. spare Annexure bullets) are dropped,
    # unless a clause edit targets them.
    edited = {int(e["paragraph_idx"]) for e in clause_edits or []}
    removed = []
    for ph in placeholders:
        p = paras[ph.paragraph_idx]
        if (
            ph.id not in values
            and ph.kind == "inline"
            and ph.paragraph_idx not in edited
            and _text(p).strip() == ph.token
            and p._p.getparent() is not None
        ):
            p._p.getparent().remove(p._p)
            removed.append(ph.id)

    applied = []
    for edit in sorted(clause_edits or [], key=lambda e: int(e["paragraph_idx"]), reverse=True):
        idx = int(edit["paragraph_idx"])
        if not 0 <= idx < len(paras):
            continue
        action = edit.get("action", "replace")
        p = paras[idx]
        if action == "delete":
            p._p.getparent().remove(p._p)
        elif action == "insert_after":
            new_el = copy.deepcopy(p._p)
            p._p.addnext(new_el)
            _set_paragraph_text(Paragraph(new_el, p._parent), edit.get("text", ""))
        else:
            _set_paragraph_text(p, edit.get("text", ""))
        applied.append({"paragraph_idx": idx, "action": action})

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    unfilled = [asdict(ph) for ph in placeholders if ph.id not in values and ph.id not in removed]
    return {
        "filled": len(values) - len(unknown),
        "unfilled": unfilled,
        "removed_empty_items": removed,
        "unknown_ids": unknown,
        "clause_edits_applied": applied,
    }


def paragraphs(path: str | Path) -> list[dict]:
    """All non-empty paragraphs with their global index and current section heading."""
    doc = Document(str(path))
    out = []
    section = ""
    for idx, p in enumerate(_paragraphs(doc)):
        text = _text(p)
        if _is_heading(p, text):
            section = text.strip()
        if text.strip():
            out.append({"paragraph_idx": idx, "section": section, "text": text.strip()})
    return out


def document_text(path: str | Path) -> str:
    return "\n".join(p["text"] for p in paragraphs(path))
