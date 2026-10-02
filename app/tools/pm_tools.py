"""Read-only Jira tools for the PM agent. Each call uses the current user's own Jira connection."""
from __future__ import annotations

import csv
import functools
from datetime import date
from typing import Optional

from langchain_core.tools import tool

from app.agents.context import require_run_user
from app.integrations.jira.client import JiraClient, JiraError
from app.integrations.jira.store import JiraNotConnected
from app.services import sprint_planner
from app.services.template_store import get_store

_last_plan: dict[int, dict] = {}


async def _client() -> JiraClient:
    return await JiraClient.for_user(require_run_user().id)


def _issue_line(i: dict) -> str:
    pts = f"{i['points']:g}pt" if i.get("points") else "–"
    extra = f" | blocked by {', '.join(i['blocked_by'])}" if i.get("blocked_by") else ""
    due = f" | due {i['due']}" if i.get("due") else ""
    return (f"{i['key']} | {i['summary'][:90]} | {i['type']} | {i['status']} | {i['priority']} | {i['assignee']} | "
            f"{pts} | sprint {i.get('sprint') or '-'}{due}{extra}")


def _errors(fn):
    @functools.wraps(fn)
    async def wrapper(*args, **kwargs):
        try:
            return await fn(*args, **kwargs)
        except JiraNotConnected as e:
            return f"JIRA NOT CONNECTED: {e} Tell the user to click 'Connect Jira' in the PM panel."
        except JiraError as e:
            return f"JIRA ERROR: {e}"
    return wrapper


@tool
@_errors
async def jira_list_projects(query: str = "") -> str:
    """List Jira projects the user can see (optionally filter by name/key text)."""
    projects = await (await _client()).projects(query)
    if not projects:
        return "No projects found."
    return f"{len(projects)} project(s):\n" + "\n".join(f"- {p['key']}: {p['name']} ({p['type']}, {p['style']})" for p in projects)


@tool
@_errors
async def jira_list_boards(project_key: str = "") -> str:
    """List boards (scrum/kanban) for a project key. Sprint and velocity tools need a scrum board id."""
    boards = await (await _client()).boards(project_key or None)
    if not boards:
        return "No boards found."
    return "\n".join(f"- id {b['id']}: {b['name']} ({b['type']}, project {b['project']})" for b in boards)


@tool
@_errors
async def jira_sprints(board_id: int, state: str = "active,future") -> str:
    """List sprints of a board. state: comma list of active, future, closed."""
    sprints = await (await _client()).sprints(board_id, state)
    if not sprints:
        return "No sprints in that state."
    return "\n".join(f"- {s['id']}: {s['name']} [{s['state']}] {(s['start'] or '')[:10]} → {(s['end'] or '')[:10]}"
                     + (f" goal: {s['goal']}" if s.get("goal") else "") for s in sprints[-25:])


@tool
@_errors
async def jira_search(jql: str, max_results: int = 50) -> str:
    """Search issues with JQL. Returns key | summary | type | status | priority | assignee | points | sprint | due | blockers."""
    issues = await (await _client()).search(jql, min(max_results, 200))
    if not issues:
        return "No issues match."
    est = sum(1 for i in issues if i.get("points"))
    pts = sum(i.get("points") or 0 for i in issues)
    head = f"{len(issues)} issue(s); {est} estimated, {pts:g} points total.\n"
    return head + "\n".join(_issue_line(i) for i in issues)


@tool
@_errors
async def jira_get_issue(key: str) -> str:
    """Full details of one issue including description."""
    i = await (await _client()).issue(key)
    return _issue_line(i) + f"\nEpic: {i.get('epic') or '-'} | Labels: {', '.join(i['labels']) or '-'}\n{i['url']}\n\n{i['description']}"


@tool
@_errors
async def jira_team(project_key: str) -> str:
    """People who can be assigned issues in a project."""
    users = await (await _client()).assignable_users(project_key)
    return "\n".join(f"- {u['name']}" for u in users if u["active"]) or "No assignable users found."


@tool
@_errors
async def jira_velocity(board_id: int, last_n: int = 5) -> str:
    """Completed story points and issue counts for the last N closed sprints of a scrum board."""
    v = await (await _client()).velocity(board_id, last_n)
    if not v["sprints"]:
        return "No closed sprints on this board, so no velocity history. Ask the PM for team capacity."
    rows = "\n".join(f"- {r['sprint']} (ended {r['end']}): {r['completed_points']:g} points, {r['completed_issues']} issues"
                     for r in v["sprints"])
    return f"{rows}\nAverage: {v['average_points']} points / {v['average_issues']} issues per sprint."


@tool(response_format="content_and_artifact")
async def plan_sprints(
    deadline: str,
    jql: str = "",
    issue_keys: Optional[list[str]] = None,
    start_date: str = "",
    sprint_length_days: int = 14,
    capacity: Optional[float] = None,
    unit: str = "auto",
    board_id: Optional[int] = None,
    buffer_pct: float = 15,
    assume_unestimated_points: Optional[float] = None,
) -> tuple[str, Optional[dict]]:
    """Build a sprint plan that fits the selected open issues before the deadline (ISO date).
    Scope: jql (e.g. 'project = QR AND statusCategory != Done ORDER BY Rank') or issue_keys.
    capacity: points (or issues) the team completes per sprint; if omitted and board_id is given, measured velocity is used.
    unit: 'points', 'issues' or 'auto' (issues when most of the scope is unestimated)."""
    try:
        client = await _client()
        if issue_keys:
            jql = f"key in ({', '.join(issue_keys)}) ORDER BY Rank"
        if not jql:
            return "ERROR: give jql or issue_keys for the scope.", None
        issues = [i for i in await client.search(jql, 300) if i.get("status_category") != "Done"]
        if not issues:
            return "No open issues in that scope.", None
        estimated_share = sum(1 for i in issues if i.get("points")) / len(issues)
        chosen_unit = unit if unit in ("points", "issues") else ("points" if estimated_share >= 0.5 else "issues")
        source = "given by user"
        if capacity is None:
            if board_id is None:
                return ("ERROR: capacity is unknown. Pass board_id to use velocity, or ask the PM for capacity "
                        f"in {chosen_unit} per sprint."), None
            v = await client.velocity(board_id, 5)
            capacity = v["average_points"] if chosen_unit == "points" else v["average_issues"]
            if not capacity:
                return f"ERROR: no velocity history in {chosen_unit}; ask the PM for capacity per sprint.", None
            source = f"average of last {len(v['sprints'])} closed sprints"
        p = sprint_planner.plan(
            issues, deadline=date.fromisoformat(deadline),
            start_date=date.fromisoformat(start_date) if start_date else date.today(),
            sprint_length_days=sprint_length_days, capacity=float(capacity), unit=chosen_unit,
            buffer_pct=buffer_pct, assume_unestimated_points=assume_unestimated_points,
        )
    except (JiraNotConnected, JiraError) as e:
        return f"JIRA ERROR: {e}", None
    except ValueError as e:
        return f"ERROR: {e}", None

    p["scope_jql"] = jql
    p["capacity_source"] = source
    p["estimated_share"] = round(estimated_share * 100)
    _last_plan[require_run_user().id] = p

    u = p["unit"]
    unit_note = ("Planning unit: ISSUE COUNT – every issue counts as 1, story points are ignored because only "
                 f"{p['estimated_share']}% of the scope is estimated." if u == "issues" else
                 "Planning unit: STORY POINTS.")
    lines = [
        unit_note,
        f"Scope: {len(issues)} open issues ({p['estimated_share']}% estimated), total {p['total_scope']:g} {u}.",
        f"Capacity: {p['capacity_per_sprint']:g} {u}/sprint ({source}), {p['buffer_pct']:g}% buffer → {p['effective_capacity']:g} usable.",
        f"Fits by {p['deadline']}: {'YES' if p['fits'] else 'NO'}.",
    ]
    for sp in p["sprints"]:
        keys = ", ".join(i["key"] for i in sp["issues"][:15]) + ("…" if len(sp["issues"]) > 15 else "")
        lines.append(f"- {sp['sprint']} {sp['start']}→{sp['end']}: {sp['planned']:g}/{sp['capacity']:g} {u} – {keys or '(empty)'}")
    if p["overflow"]:
        lines.append(f"Overflow: {len(p['overflow'])} issues / {p['overflow_size']:g} {u} won't fit "
                     f"({', '.join(i['key'] for i in p['overflow'][:15])}). Needs {p['extra_sprints_needed']} more sprint(s) "
                     f"→ projected completion {p['projected_completion']}, or ~{p['capacity_needed_per_sprint']} {u}/sprint capacity.")
    if p["risks"]:
        lines.append("Risks: " + " | ".join(p["risks"]))
    cols, rows = sprint_planner.plan_rows(p)
    return "\n".join(lines), {"type": "table", "title": f"Sprint plan to {p['deadline']}", "columns": cols, "rows": rows}


@tool(response_format="content_and_artifact")
def export_plan(format: str = "xlsx") -> tuple[str, Optional[dict]]:
    """Export the most recent sprint plan as a downloadable file: xlsx, csv or md."""
    p = _last_plan.get(require_run_user().id)
    if not p:
        return "ERROR: no plan yet; run plan_sprints first.", None
    fmt = format.lower().strip(".")
    if fmt not in ("xlsx", "csv", "md"):
        return "ERROR: format must be xlsx, csv or md.", None
    cols, rows = sprint_planner.plan_rows(p)
    store = get_store()
    file_id, _ = store.new_output_path(f"sprint_plan_{p['deadline']}")
    path = store.outputs_dir / f"{file_id}.{fmt}"
    if fmt == "csv":
        with path.open("w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(cols)
            w.writerows(rows)
    elif fmt == "md":
        md = [f"# Sprint plan to {p['deadline']}", "", f"Capacity {p['capacity_per_sprint']} {p['unit']}/sprint, "
              f"fits: {'yes' if p['fits'] else 'no'}", "", "| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
        md += ["| " + " | ".join("" if v is None else str(v) for v in r) + " |" for r in rows]
        if p["risks"]:
            md += ["", "## Risks", *[f"- {r}" for r in p["risks"]]]
        path.write_text("\n".join(md))
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Plan"
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="7B68E2")
        for r in rows:
            ws.append(r)
        for col, width in zip("ABCDEFGHI", (24, 12, 12, 10, 60, 8, 10, 22, 14)):
            ws.column_dimensions[col].width = width
        summary = wb.create_sheet("Summary")
        for k in ("deadline", "start_date", "unit", "capacity_per_sprint", "effective_capacity", "total_scope", "fits",
                  "overflow_size", "extra_sprints_needed", "projected_completion", "scope_jql", "capacity_source"):
            summary.append([k, str(p.get(k))])
        for r in p["risks"]:
            summary.append(["risk", r])
        summary.column_dimensions["A"].width = 24
        summary.column_dimensions["B"].width = 100
        wb.save(path)
    mime = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "csv": "text/csv",
            "md": "text/markdown"}[fmt]
    return f"Exported plan as {path.name}.", {"type": "file", "file_id": file_id, "filename": path.name, "mime": mime}


PM_TOOLS = [jira_list_projects, jira_list_boards, jira_sprints, jira_search, jira_get_issue, jira_team,
            jira_velocity, plan_sprints, export_plan]
