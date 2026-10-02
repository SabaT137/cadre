from pathlib import Path

from docx import Document

from app.services import docx_filler

TEMPLATE = Path(__file__).resolve().parent.parent / "storage" / "templates" / "consultant_hourly_parttime.docx"


def _all_text(path) -> str:
    return docx_filler.document_text(path)


def test_scan_finds_all_bullet_placeholders():
    phs = docx_filler.scan(TEMPLATE)
    inline = [p for p in phs if p["kind"] == "inline"]
    raw = _all_text(TEMPLATE)
    assert len(inline) == raw.count("[●]") + raw.count("[insert relevant details]")
    assert any("CNIC" in p["context"] for p in phs)
    assert sum(1 for p in phs if p["section"].startswith("ANNEXURE")) == 12
    assert any(p["kind"] == "blank_field" for p in phs)


def test_fill_all_replaces_every_placeholder(tmp_path):
    phs = docx_filler.scan(TEMPLATE)
    out = tmp_path / "out.docx"
    result = docx_filler.fill(TEMPLATE, out, {str(p["id"]): f"VAL{p['id']}" for p in phs})
    text = _all_text(out)
    assert "[●]" not in text and "[insert relevant details]" not in text
    assert result["unfilled"] == [] and result["unknown_ids"] == []
    assert "holding CNIC No. VAL5 resident of VAL6" in text
    assert "Name: VAL12" in text  # blank field gets the value appended


def test_split_run_placeholder_and_formatting_preserved(tmp_path):
    # Placeholder #10 ("two [●]  year(s)") is split across runs "[" and "●]" in the template.
    out = tmp_path / "out.docx"
    docx_filler.fill(TEMPLATE, out, {"10": "1"})
    assert "successive periods of two 1  year(s)" in _all_text(out)
    # Bold run "Agreement" in the first paragraph stays bold after filling the same paragraph.
    docx_filler.fill(TEMPLATE, out, {"1": "2nd", "2": "October", "3": "2026"})
    p = next(p for p in Document(str(out)).paragraphs if "Effective Date" in p.text)
    assert "on the 2nd day of October in the year 2026" in p.text
    original = next(p for p in Document(str(TEMPLATE)).paragraphs if "Effective Date" in p.text)
    assert [r.bold for r in p.runs] == [r.bold for r in original.runs]


def test_multiline_bullet_value_becomes_separate_bullets(tmp_path):
    out = tmp_path / "out.docx"
    first = next(p for p in docx_filler.scan(TEMPLATE) if p["section"].startswith("ANNEXURE"))
    docx_filler.fill(TEMPLATE, out, {str(first["id"]): "- Design screens\n- Run user tests\n- Keep design system"})
    paras = [p["text"] for p in docx_filler.paragraphs(out) if p["section"].startswith("ANNEXURE")]
    assert paras[-3:] == ["Design screens", "Run user tests", "Keep design system"]


def test_unused_bullets_removed_and_clause_edits(tmp_path):
    out = tmp_path / "out.docx"
    phs = docx_filler.scan(TEMPLATE)
    bullets = [p for p in phs if p["section"].startswith("ANNEXURE")]
    notice = next(p for p in docx_filler.paragraphs(TEMPLATE) if "7 days" in p["text"])
    result = docx_filler.fill(
        TEMPLATE, out,
        {str(bullets[0]["id"]): "Build APIs", str(bullets[1]["id"]): "Review code"},
        [{"paragraph_idx": notice["paragraph_idx"], "action": "replace",
          "text": "Either Party may terminate this Agreement at any time, by providing 30 days' written notice."}],
    )
    text = _all_text(out)
    assert len(result["removed_empty_items"]) == 10
    assert "Build APIs" in text and "Review code" in text
    assert "30 days' written notice" in text and "7 days" not in text
